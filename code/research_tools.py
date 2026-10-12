"""Local research tools for Codex: real sources, page provenance, editable manuscripts.

Independently implemented adapters. Interpretation and drafting are performed by
Codex; these tools retrieve records and operate on user-selected files.
"""
from __future__ import annotations

import csv
from contextlib import closing
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote, unquote

import requests


def _path(workspace: str, relative: str, exists: bool = False) -> Path:
    root = Path(workspace).expanduser().resolve(strict=True)
    if not root.is_dir():
        raise ValueError('workspace must be an existing directory')
    if not relative or Path(relative).is_absolute() or ':' in relative:
        raise ValueError('Use a relative path within the selected workspace')
    path = (root / relative).resolve()
    if not path.is_relative_to(root) or path == root:
        raise ValueError('Path escapes the selected workspace')
    if exists and not path.is_file():
        raise FileNotFoundError(relative)
    return path


def _write(path: Path, data: str, overwrite: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w' if overwrite else 'x', encoding='utf-8', newline='') as handle:
        handle.write(data)


def _doi(value: str) -> str:
    value = unquote(value.strip())
    value = re.sub(r'^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)', '', value, flags=re.I)
    if not re.fullmatch(r'10\.\d{4,9}/[^\s]+', value):
        raise ValueError('Expected a DOI, such as 10.1038/nature12373')
    return value.lower()


def _crossref(endpoint: str, params: dict | None = None) -> dict:
    response = requests.get('https://api.crossref.org/works' + endpoint,
                            params=params, timeout=(10, 45),
                            headers={'User-Agent': 'CodexScience/0.2 (research metadata adapter)'})
    response.raise_for_status()
    payload = response.json()
    if payload.get('status') != 'ok' or not isinstance(payload.get('message'), dict):
        raise ValueError('Unexpected Crossref response')
    return payload['message']


def _record(item: dict, detailed: bool = True) -> dict:
    dates = item.get('published', item.get('issued', {})).get('date-parts', [[]])
    authors = [{'given': a.get('given', ''), 'family': a.get('family', ''),
                'name': a.get('name', '')} for a in item.get('author', [])]
    doi = _doi(item['DOI'])
    return {'doi': doi, 'title': ' '.join(item.get('title', [])), 'authors': authors,
            'year': dates[0][0] if dates and dates[0] else None,
            'journal': ' '.join(item.get('container-title', [])),
            'volume': item.get('volume', ''), 'issue': item.get('issue', ''),
            'pages': item.get('page', ''), 'type': item.get('type', ''),
            'url': 'https://doi.org/' + quote(doi, safe='/'),
            'abstract': str(item.get('abstract', ''))[:20000] if detailed else None,
            'references': item.get('reference', [])[:200] if detailed else [],
            'reference_count': len(item.get('reference', [])),
            'references_truncated': len(item.get('reference', [])) > (200 if detailed else 0),
            'updates': item.get('update-to', []), 'relations': item.get('relation', {}),
            'source': 'Crossref', 'retrieved_at': datetime.now(timezone.utc).isoformat(),
            'verification_scope': 'Bibliographic metadata; claim support and retraction status require further review.'}


def science_crossref_search(query: str, limit: int = 10, work_type: str = 'journal-article') -> dict:
    """Search actual Crossref paper metadata. Results are candidates; inspect titles before citing."""
    if not query.strip() or not 1 <= limit <= 50:
        raise ValueError('Provide a query and limit between 1 and 50')
    types = {'journal-article', 'posted-content', 'proceedings-article', 'book-chapter', 'book', 'all'}
    if work_type not in types:
        raise ValueError('work_type must be journal-article, posted-content, proceedings-article, book-chapter, book, or all')
    params = {'query.bibliographic': query, 'rows': limit}
    if work_type != 'all':
        params['filter'] = 'type:' + work_type
    message = _crossref('', params)
    return {'query': query, 'total_results': message.get('total-results'),
            'work_type': work_type, 'records': [_record(item, detailed=False) for item in message.get('items', [])]}


def science_doi_lookup(doi: str) -> dict:
    """Resolve a DOI to Crossref metadata, references, and available update notices."""
    normalized = _doi(doi)
    record = _record(_crossref('/' + quote(normalized, safe='')))
    if record['doi'] != normalized:
        raise ValueError('Returned DOI differs from requested DOI')
    return record


def _line(value: object) -> str:
    return re.sub(r'[\r\n\t]+', ' ', str(value or '')).strip()


def _bib(value: object) -> str:
    text = _line(value)
    return ''.join({'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}',
                    '%': r'\%', '&': r'\&', '#': r'\#', '_': r'\_',
                    '$': r'\$', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}.get(c, c) for c in text)


def science_export_bibliography(workspace: str, dois: list[str], output: str = 'references.bib',
                                format: str = 'bibtex', overwrite: bool = False) -> dict:
    """Retrieve each DOI and export a BibTeX, RIS, or JSON bibliography; fail if a DOI cannot resolve."""
    target = _path(workspace, output)
    if format not in {'bibtex', 'ris', 'json'} or not 1 <= len(dois) <= 100:
        raise ValueError('format must be bibtex, ris, or json; provide 1–100 DOIs')
    records = [science_doi_lookup(doi) for doi in dict.fromkeys(_doi(d) for d in dois)]
    if format == 'json':
        body = json.dumps(records, ensure_ascii=False, indent=2)
    elif format == 'ris':
        entries = []
        for record in records:
            lines = ['TY  - ' + ('JOUR' if record['type'] == 'journal-article' else 'GEN'),
                     'TI  - ' + _line(record['title'])]
            for author in record['authors']:
                name = author['name'] or ', '.join(filter(None, [author['family'], author['given']]))
                lines.append('AU  - ' + _line(name))
            for tag, key in [('PY', 'year'), ('JO', 'journal'), ('VL', 'volume'), ('IS', 'issue'), ('DO', 'doi'), ('UR', 'url')]:
                if record[key]:
                    lines.append(f'{tag}  - {_line(record[key])}')
            lines.append('ER  - ')
            entries.append('\n'.join(lines))
        body = '\n\n'.join(entries) + '\n'
    else:
        entries = []
        for record in records:
            key = 'doi_' + hashlib.sha256(record['doi'].encode()).hexdigest()[:12]
            authors = []
            for author in record['authors']:
                authors.append('{' + _bib(author['name']) + '}' if author['name'] else
                               ', '.join(_bib(v) for v in [author['family'], author['given']] if v))
            fields = {'title': _bib(record['title']), 'author': ' and '.join(authors),
                      'year': _bib(record['year']), 'journal': _bib(record['journal']),
                      'volume': _bib(record['volume']), 'number': _bib(record['issue']),
                      'pages': _bib(record['pages']), 'doi': _bib(record['doi']), 'url': _bib(record['url'])}
            kind = 'article' if record['type'] == 'journal-article' else 'misc'
            entries.append('@' + kind + '{' + key + ',\n' +
                           '\n'.join(f'  {field} = {{{value}}},' for field, value in fields.items() if value) + '\n}')
        body = '\n\n'.join(entries) + '\n'
    _write(target, body, overwrite)
    compact = [{key: record[key] for key in ('doi', 'title', 'year', 'journal', 'url')} for record in records]
    return {'path': str(target), 'count': len(records), 'format': format, 'records': compact}


def _pdf(workspace: str, path: str):
    import pypdfium2
    return pypdfium2.PdfDocument(str(_path(workspace, path, exists=True)))


def science_pdf_pages(workspace: str, path: str, pages: list[int] | None = None,
                      max_chars_per_page: int = 16000) -> dict:
    """Read PDF pages with provenance, capped at 200000 text characters per response. Pass next_pages to continue; image-only PDFs need OCR separately."""
    if not 100 <= max_chars_per_page <= 100000:
        raise ValueError('max_chars_per_page must be between 100 and 100000')
    with _pdf(workspace, path) as document:
        selected = pages if pages is not None else list(range(1, min(len(document), 20) + 1))
        if len(selected) > 100 or any(not isinstance(p, int) or not 1 <= p <= len(document) for p in selected):
            raise ValueError('Select up to 100 valid one-based page numbers')
        rows = []
        total_chars = 0
        for number in selected:
            # Reserve a full per-page allowance so continuation always begins on a whole page.
            if 200000 - total_chars < max_chars_per_page:
                break
            with closing(document[number - 1]) as page:
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_range()
            rows.append({'page': number, 'text': text[:max_chars_per_page],
                         'truncated': len(text) > max_chars_per_page, 'characters': len(text),
                         'has_text': bool(text.strip()), 'source': f'{path}#page={number}'})
            total_chars += min(len(text), max_chars_per_page)
        pending = selected[len(rows):]
        next_pages = pending
        if pages is None and not pending and selected:
            next_pages = list(range(selected[-1] + 1, min(len(document), selected[-1] + 20) + 1))
        return {'path': path, 'page_count': len(document), 'pages': rows,
                'returned_characters': total_chars, 'text_character_budget': 200000,
                'pages_read': len(rows), 'remaining_requested_pages': pending, 'next_pages': next_pages,
                'remaining_pages': len(document) - len({row['page'] for row in rows})}


def science_pdf_outline(workspace: str, path: str) -> dict:
    """Read PDF bookmark headings and destinations; a PDF may have no embedded outline."""
    with _pdf(workspace, path) as document:
        rows = []
        for bookmark in document.get_toc():
            destination = bookmark.get_dest()
            index = destination.get_index() if destination else None
            rows.append({'heading': bookmark.get_title(), 'level': bookmark.level,
                         'page': index + 1 if index is not None and index >= 0 else None})
        return {'path': path, 'page_count': len(document), 'outline': rows}


def science_pdf_search(workspace: str, path: str, query: str, top_k: int = 10) -> dict:
    """Rank PDF pages by literal query-term matches and return excerpts; this is lexical search."""
    terms = list(dict.fromkeys(re.findall(r'\w+', query.casefold())))
    if not terms or not 1 <= top_k <= 50:
        raise ValueError('Provide search terms and top_k between 1 and 50')
    hits = []
    with _pdf(workspace, path) as document:
        for index in range(len(document)):
            with closing(document[index]) as page:
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_range()
            folded = text.casefold()
            counts = {term: folded.count(term) for term in terms}
            score = sum(counts.values())
            if score:
                start = max(0, min(folded.find(term) for term in terms if counts[term]) - 160)
                hits.append({'page': index + 1, 'score': score, 'matches': counts,
                             'excerpt': text[start:start + 1200], 'source': f'{path}#page={index + 1}'})
        hits.sort(key=lambda row: (-row['score'], row['page']))
        return {'path': path, 'method': 'literal-term-frequency', 'pages_scanned': len(document),
                'matched_pages': len(hits), 'hits': hits[:top_k]}


def science_pdf_extract_identifiers(workspace: str, path: str) -> dict:
    """Extract DOI and GEO accession candidates from every PDF text page, retaining page locations."""
    patterns = {'doi': re.compile(r'10\.\d{4,9}/[^\s<>\"]+', re.I),
                'geo': re.compile(r'\bG(?:SE|SM|PL)\d+\b')}
    found = {}
    with _pdf(workspace, path) as document:
        for index in range(len(document)):
            with closing(document[index]) as page:
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_range()
            for kind, pattern in patterns.items():
                for match in pattern.finditer(text):
                    value = match.group().rstrip('.,;')
                    key = (kind, value)
                    row = found.setdefault(key, {'kind': kind, 'value': value, 'pages': [], 'verified': False})
                    if index + 1 not in row['pages']:
                        row['pages'].append(index + 1)
        return {'path': path, 'pages_scanned': len(document), 'candidates': list(found.values()),
                'note': 'PDF line breaks or punctuation may affect candidates; resolve DOIs before citing.'}


def science_pdf_render_page(workspace: str, path: str, page: int, output: str,
                            dpi: int = 150, overwrite: bool = False) -> dict:
    """Render one PDF page to a PNG for figure/table inspection with Codex image viewing."""
    target = _path(workspace, output)
    if target.suffix.lower() != '.png' or not 72 <= dpi <= 300:
        raise ValueError('Choose a .png output and dpi between 72 and 300')
    if target.exists() and not overwrite:
        raise FileExistsError(str(target))
    with _pdf(workspace, path) as document:
        if not 1 <= page <= len(document):
            raise ValueError('Invalid one-based page number')
        with closing(document[page - 1]) as selected:
            width, height = selected.get_size()
            if width * height * (dpi / 72) ** 2 > 50_000_000:
                raise ValueError('Rendered page would exceed 50 megapixels; reduce dpi')
            bitmap = selected.render(scale=dpi / 72)
            try:
                image = bitmap.to_pil()
                target.parent.mkdir(parents=True, exist_ok=True)
                with target.open('wb' if overwrite else 'xb') as handle:
                    image.save(handle, format='PNG')
            finally:
                bitmap.close()
    return {'path': str(target), 'source': f'{path}#page={page}', 'dpi': dpi}


def science_save_evidence(workspace: str, entries: list[dict], output: str = 'evidence.csv',
                          overwrite: bool = False) -> dict:
    """Save a supplied evidence table. Local PDF quotes are checked against their cited page; claims remain author interpretation."""
    target = _path(workspace, output)
    if not 1 <= len(entries) <= 1000:
        raise ValueError('Provide 1–1000 evidence entries')
    columns = ['claim', 'source', 'page', 'quote', 'study_design', 'finding', 'limitations', 'quote_check']
    rows = []
    for entry in entries:
        if not str(entry.get('claim', '')).strip() or not str(entry.get('source', '')).strip():
            raise ValueError('Every entry needs a claim and source')
        row = {key: entry.get(key, '') for key in columns}
        row['quote_check'] = 'not_checked'
        if entry.get('quote') and not str(entry['quote']).strip():
            raise ValueError('Quotation must contain non-whitespace text')
        if str(entry['source']).lower().endswith('.pdf') and entry.get('quote'):
            if not isinstance(entry.get('page'), int):
                raise ValueError('PDF quotes need an integer page number')
            page = science_pdf_pages(workspace, entry['source'], [entry['page']], 100000)['pages'][0]
            normalize = lambda value: ' '.join(str(value).split())
            if normalize(entry['quote']) not in normalize(page['text']):
                raise ValueError(f'Quote not found on cited page: {entry["source"]}#{entry["page"]}')
            row['quote_check'] = 'matched_pdf_text'
        rows.append(row)
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream, fieldnames=columns)
    writer.writeheader()
    writer.writerows(rows)
    _write(target, stream.getvalue(), overwrite)
    return {'path': str(target), 'entries': len(rows), 'quote_checks': [row['quote_check'] for row in rows],
            'scope': 'Quote matching checks text presence; study interpretation and claim support require review.'}


def science_create_manuscript(workspace: str, title: str, output: str = 'manuscript.md',
                               sections: list[str] | None = None) -> dict:
    """Create an editable manuscript scaffold with empty sections; never invent study findings."""
    if not title.strip():
        raise ValueError('A manuscript title is required')
    headings = sections or ['Abstract', 'Introduction', 'Methods', 'Results', 'Discussion', 'References']
    body = '# ' + _line(title) + '\n\n' + '\n\n'.join('## ' + _line(s) + '\n' for s in headings) + '\n'
    target = _path(workspace, output)
    _write(target, body)
    return {'path': str(target), 'sections': headings}


_IMAGE_PATTERN = re.compile(r'!\[([^\]]*)\]\((<[^>]+>|[^)]+)\)')


def _image_path(workspace: str, original: Path, reference: str) -> Path:
    value = reference.strip()
    if value.startswith('<') and value.endswith('>'):
        value = value[1:-1]
    if re.match(r'^[a-zA-Z][a-zA-Z0-9+.-]*:', value) or value.startswith('//'):
        raise ValueError('Figure images must be local paths within the workspace')
    root = Path(workspace).expanduser().resolve(strict=True)
    relative = original.parent.relative_to(root) / unquote(value)
    return _path(workspace, str(relative), exists=True)


def _append_inline(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    pattern = re.compile(r'`([^`]+)`|\[([^\]]+)\]\((<[^>]+>|[^)]+)\)|\*\*(.+?)\*\*|(?<!\*)\*([^*]+)\*(?!\*)')
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold, run.italic = bold, italic
        if match[1] is not None:
            run = paragraph.add_run(match[1])
            run.font.name = 'Consolas'
            run.bold, run.italic = bold, italic
        elif match[2] is not None:
            url = match[3].strip('<>')
            if re.match(r'^(https?://|mailto:)', url, re.I):
                hyperlink = OxmlElement('w:hyperlink')
                hyperlink.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
                run = OxmlElement('w:r')
                properties = OxmlElement('w:rPr')
                style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink')
                properties.append(style)
                if bold:
                    properties.append(OxmlElement('w:b'))
                if italic:
                    properties.append(OxmlElement('w:i'))
                run.append(properties)
                content = OxmlElement('w:t'); content.text = match[2]
                run.append(content); hyperlink.append(run); paragraph._p.append(hyperlink)
            else:
                run = paragraph.add_run(match.group())
                run.bold, run.italic = bold, italic
        elif match[4] is not None:
            _append_inline(paragraph, match[4], bold=True, italic=italic)
        else:
            _append_inline(paragraph, match[5], bold=bold, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold, run.italic = bold, italic


def _table_cells(line: str) -> list[str]:
    return [cell.strip().replace(r'\|', '|') for cell in re.split(r'(?<!\\)\|', line.strip().strip('|'))]


def science_check_manuscript(workspace: str, source: str) -> dict:
    """Find empty sections, unresolved citation placeholders and missing local figures; this checks document structure, not scientific validity."""
    original = _path(workspace, source, exists=True)
    lines = original.read_text(encoding='utf-8-sig').splitlines()
    findings = []
    headings = []
    references = set()
    cited_numbers = []
    doi_candidates = []
    figures = []
    in_code = False
    in_references = False
    visible = []
    for number, line in enumerate(lines, 1):
        if line.lstrip().startswith('```'):
            in_code = not in_code
            visible.append('')
            continue
        if in_code:
            visible.append('')
            continue
        visible.append(line)
        heading = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading:
            headings.append((number, len(heading[1]), heading[2]))
            in_references = heading[2].strip().casefold() in {'references', 'bibliography', '参考文献'}
        definition = re.match(r'^\s*\[(\d+)\]:' if not in_references else r'^\s*\[(\d+)\](?::|\s+)', line)
        if definition:
            references.add(definition[1])
        for match in re.finditer(r'\b(?:TODO|TBD|FIXME)\b|\[(?:citation needed|ref(?:erence)?\s*\??|doi\s*[:?]?|\?+)\]|\{\{[^}]+\}\}', line, re.I):
            findings.append({'kind': 'unresolved_placeholder', 'line': number, 'text': match.group()})
        for match in re.finditer(r'(?<!!)\[(\d+)\](?![:(])', line):
            if not definition or match.start() != definition.start():
                cited_numbers.append((number, match[1]))
        for match in re.finditer(r'(?:https?://(?:dx\.)?doi\.org/|\bDOI\s*:\s*)([^\s<>]+)', line, re.I):
            candidate = match[1].rstrip('.,;')
            while candidate.endswith(')') and candidate.count(')') > candidate.count('('):
                candidate = candidate[:-1]
            try:
                normalized = _doi(candidate)
            except ValueError:
                findings.append({'kind': 'malformed_doi', 'line': number, 'text': candidate})
            else:
                doi_candidates.append({'doi': normalized, 'line': number, 'resolved_online': False})
        for image in _IMAGE_PATTERN.finditer(line):
            try:
                figure = _image_path(workspace, original, image[2])
            except (FileNotFoundError, ValueError) as error:
                findings.append({'kind': 'missing_or_invalid_figure', 'line': number,
                                 'path': image[2], 'detail': str(error)})
            else:
                figures.append({'line': number, 'path': str(figure), 'caption': image[1]})
    for index, (number, level, title) in enumerate(headings):
        stop = next((n for n, depth, _ in headings[index + 1:] if depth <= level), len(lines) + 1)
        body = [line.strip() for line in visible[number:stop - 1] if line.strip() and not line.lstrip().startswith('#')]
        if not body:
            findings.append({'kind': 'empty_section', 'line': number, 'heading': title})
    for number, label in cited_numbers:
        if label not in references:
            findings.append({'kind': 'unresolved_numbered_reference', 'line': number, 'reference': label})
    if in_code:
        findings.append({'kind': 'unclosed_code_fence', 'line': len(lines)})
    return {'source': source, 'findings': findings, 'figures': figures,
            'doi_candidates': doi_candidates, 'section_count': len(headings),
            'scope': 'Local structural checks. Resolve DOI candidates with science_doi_lookup; review claim support against original sources.'}


def science_export_manuscript(workspace: str, source: str, output: str,
                               overwrite: bool = False) -> dict:
    """Export Markdown to DOCX with headings, tables, emphasis, hyperlinks and local figure images. Captions are editable paragraphs; equations remain literal."""
    from docx import Document
    from docx.shared import Inches
    original = _path(workspace, source, exists=True)
    target = _path(workspace, output)
    if target.suffix.lower() != '.docx':
        raise ValueError('Output must end in .docx')
    if original == target:
        raise ValueError('Source and output must differ')
    if target.exists() and not overwrite:
        raise FileExistsError(str(target))
    document = Document()
    paragraph = []
    in_code = False
    warnings = []
    def flush():
        if paragraph:
            _append_inline(document.add_paragraph(), ' '.join(paragraph))
            paragraph.clear()
    lines = original.read_text(encoding='utf-8-sig').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith('```'):
            flush()
            in_code = not in_code
        elif in_code:
            run = document.add_paragraph(style='No Spacing').add_run(line)
            run.font.name = 'Consolas'
        elif not line.strip():
            flush()
        elif match := re.match(r'^(#{1,6})\s+(.*)$', line):
            flush()
            _append_inline(document.add_heading(level=len(match[1])), match[2])
        elif image := _IMAGE_PATTERN.fullmatch(line.strip()):
            flush()
            figure = _image_path(workspace, original, image[2])
            document.add_picture(str(figure), width=Inches(6))
            if image[1].strip():
                _append_inline(document.add_paragraph(), image[1])
        elif index + 1 < len(lines) and '|' in line and all(re.fullmatch(r':?-{3,}:?', c) for c in _table_cells(lines[index + 1])):
            flush()
            header = _table_cells(line)
            separators = _table_cells(lines[index + 1])
            if len(header) != len(separators):
                raise ValueError(f'Table header and delimiter column counts differ at line {index + 1}')
            table = document.add_table(rows=1, cols=len(header)); table.style = 'Table Grid'
            for cell, value in zip(table.rows[0].cells, header):
                _append_inline(cell.paragraphs[0], value, bold=True)
            index += 2
            while index < len(lines) and '|' in lines[index] and lines[index].strip():
                values = _table_cells(lines[index])
                if len(values) != len(header):
                    raise ValueError(f'Table column count differs at line {index + 1}')
                for cell, value in zip(table.add_row().cells, values):
                    _append_inline(cell.paragraphs[0], value)
                index += 1
            continue
        elif re.match(r'^[-*]\s+', line):
            flush()
            _append_inline(document.add_paragraph(style='List Bullet'), line[2:])
        elif match := re.match(r'^\d+\.\s+(.*)$', line):
            flush()
            _append_inline(document.add_paragraph(style='List Number'), match[1])
        else:
            if _IMAGE_PATTERN.search(line):
                warnings.append(f'Inline image at line {index + 1} retained as text; place images on separate lines to embed them.')
            paragraph.append(line)
        index += 1
    flush()
    if in_code:
        warnings.append('Source ends inside a fenced code block.')
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open('wb' if overwrite else 'xb') as handle:
        document.save(handle)
    return {'path': str(target), 'source': source, 'paragraphs': len(document.paragraphs),
            'tables': len(document.tables), 'figures': len(document.inline_shapes), 'warnings': warnings,
            'format_scope': 'Headings, prose, lists, code, simple pipe tables, emphasis, HTTP/mailto links and standalone local figures. Equations and unsupported markup remain literal text.'}


TOOLS = [science_crossref_search, science_doi_lookup, science_export_bibliography,
         science_pdf_pages, science_pdf_outline, science_pdf_search,
         science_pdf_extract_identifiers, science_pdf_render_page, science_save_evidence,
         science_create_manuscript, science_check_manuscript, science_export_manuscript]


def register_tools(mcp) -> None:
    for function in TOOLS:
        mcp.tool()(function)
