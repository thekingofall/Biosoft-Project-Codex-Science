"""Behavior checks for research tools using a generated two-page PDF and mocked DOI records."""
import csv
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'code'))
import research_tools as rt


def make_pdf(path):
    texts = ['Methods: BRCA1 cohort. DOI 10.1038/nature12373 GEO GSE123.',
             'Results: TP53 was measured. Limitations include sample size.']
    objects = [b'<< /Type /Catalog /Pages 2 0 R /Outlines 8 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>']
    for page_id, text in [(3, texts[0]), (5, texts[1])]:
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents {page_id + 1} 0 R >>'.encode())
        stream = ('BT /F1 10 Tf 20 740 Td (' + text + ') Tj ET').encode()
        objects.append(f'<< /Length {len(stream)} >>\nstream\n'.encode() + stream + b'\nendstream')
    objects.extend([b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
                    b'<< /Type /Outlines /First 9 0 R /Last 9 0 R /Count 1 >>',
                    b'<< /Title (Methods) /Parent 8 0 R /Dest [3 0 R /Fit] >>'])
    data = b'%PDF-1.4\n'
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f'{index} 0 obj\n'.encode() + obj + b'\nendobj\n'
    xref = len(data)
    data += f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode()
    data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    data += f'trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    path.write_bytes(data)


RECORD = {'DOI': '10.1038/nature12373', 'title': ['A & B: {Research}'],
          'author': [{'given': 'Alice', 'family': 'Smith'}],
          'published': {'date-parts': [[2013, 9, 1]]},
          'container-title': ['Nature'], 'type': 'journal-article',
          'reference': [{'DOI': '10.1000/example'}]}


class ResearchTests(unittest.TestCase):
    def setUp(self):
        parent = os.environ.get('CODEX_SCIENCE_TEST_TMP')
        if parent:
            Path(parent).mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=parent)
        self.root = Path(self.temp.name)
        self.workspace = str(self.root)
        make_pdf(self.root / 'paper.pdf')

    def tearDown(self):
        self.temp.cleanup()

    def test_paths_and_existing_files(self):
        for relative in ['../escape.md', self.workspace, 'C:escape', '']:
            with self.assertRaises(ValueError):
                rt.science_create_manuscript(self.workspace, 'Study', relative)
        first = rt.science_create_manuscript(self.workspace, 'Study')
        with self.assertRaises(FileExistsError):
            rt.science_create_manuscript(self.workspace, 'Changed')
        self.assertIn('# Study', Path(first['path']).read_text())

    def test_doi_metadata_and_mismatch(self):
        with patch.object(rt, '_crossref', return_value=RECORD):
            result = rt.science_doi_lookup('https://doi.org/10.1038/NATURE12373')
            self.assertEqual(result['year'], 2013)
            self.assertEqual(result['authors'][0]['family'], 'Smith')
            with self.assertRaises(ValueError):
                rt.science_doi_lookup('10.1000/other')
        with self.assertRaises(ValueError):
            rt.science_doi_lookup('invented citation')

    def test_search_candidates_and_bounds(self):
        with patch.object(rt, '_crossref', return_value={'items': [RECORD], 'total-results': 1}):
            result = rt.science_crossref_search('Research', 1)
            self.assertEqual(len(result['records']), 1)
        with self.assertRaises(ValueError):
            rt.science_crossref_search('Research', 1000)

    def test_bibliography_retrieval_dedup_and_escaping(self):
        with patch.object(rt, '_crossref', return_value=RECORD) as lookup:
            result = rt.science_export_bibliography(self.workspace,
                ['10.1038/nature12373', 'https://doi.org/10.1038/nature12373'])
            self.assertEqual(result['count'], 1)
            self.assertEqual(lookup.call_count, 1)
            text = (self.root / 'references.bib').read_text()
            self.assertIn(r'A \& B: \{Research\}', text)
            self.assertIn('Smith, Alice', text)
            rt.science_export_bibliography(self.workspace, ['10.1038/nature12373'], 'refs.ris', 'ris')
            self.assertIn('DO  - 10.1038/nature12373', (self.root / 'refs.ris').read_text())
            rt.science_export_bibliography(self.workspace, ['10.1038/nature12373'], 'refs.json', 'json')
            self.assertEqual(json.loads((self.root / 'refs.json').read_text())[0]['year'], 2013)

    def test_lookup_failure_leaves_no_export(self):
        with patch.object(rt, '_crossref', side_effect=ValueError('not found')):
            with self.assertRaises(ValueError):
                rt.science_export_bibliography(self.workspace, ['10.1000/missing'])
        self.assertFalse((self.root / 'references.bib').exists())

    @unittest.skipUnless(importlib.util.find_spec('pypdfium2'), 'pypdfium2 unavailable')
    def test_pdf_pages_outline_and_search(self):
        result = rt.science_pdf_pages(self.workspace, 'paper.pdf', [2, 1])
        self.assertEqual(result['page_count'], 2)
        self.assertIn('TP53', result['pages'][0]['text'])
        self.assertEqual(result['pages'][0]['source'], 'paper.pdf#page=2')
        self.assertEqual(rt.science_pdf_outline(self.workspace, 'paper.pdf')['outline'][0]['page'], 1)
        hits = rt.science_pdf_search(self.workspace, 'paper.pdf', 'BRCA1')['hits']
        self.assertEqual([h['page'] for h in hits], [1])
        self.assertEqual(rt.science_pdf_search(self.workspace, 'paper.pdf', 'nonexistent')['hits'], [])
        with self.assertRaises(ValueError):
            rt.science_pdf_pages(self.workspace, 'paper.pdf', [3])

    @unittest.skipUnless(importlib.util.find_spec('pypdfium2'), 'pypdfium2 unavailable')
    def test_pdf_identifiers_and_quote_checks(self):
        identifiers = rt.science_pdf_extract_identifiers(self.workspace, 'paper.pdf')['candidates']
        self.assertIn('10.1038/nature12373', [r['value'] for r in identifiers])
        self.assertIn('GSE123', [r['value'] for r in identifiers])
        entry = {'claim': 'The study includes a BRCA1 cohort', 'source': 'paper.pdf',
                 'page': 1, 'quote': 'BRCA1 cohort'}
        result = rt.science_save_evidence(self.workspace, [entry])
        self.assertEqual(result['quote_checks'], ['matched_pdf_text'])
        with (self.root / 'evidence.csv').open(newline='') as stream:
            self.assertEqual(list(csv.DictReader(stream))[0]['page'], '1')
        entry['quote'] = 'A fabricated result'
        with self.assertRaises(ValueError):
            rt.science_save_evidence(self.workspace, [entry], 'invalid.csv')
        self.assertFalse((self.root / 'invalid.csv').exists())

    @unittest.skipUnless(importlib.util.find_spec('pypdfium2') and importlib.util.find_spec('PIL'), 'PDF rendering dependencies unavailable')
    def test_render_real_pdf_page(self):
        from PIL import Image
        result = rt.science_pdf_render_page(self.workspace, 'paper.pdf', 1, 'figures/page.png', 72)
        with Image.open(result['path']) as image:
            self.assertEqual(image.size, (612, 792))
        with self.assertRaises(FileExistsError):
            rt.science_pdf_render_page(self.workspace, 'paper.pdf', 1, 'figures/page.png')

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'python-docx unavailable')
    def test_docx_round_trip(self):
        from docx import Document
        (self.root / 'draft.md').write_text('# Real Study\n\n## Methods\n\nMeasured BRCA1.\n\n- Cohort A\n- Cohort B\n', encoding='utf-8')
        result = rt.science_export_manuscript(self.workspace, 'draft.md', 'draft.docx')
        document = Document(result['path'])
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('Measured BRCA1.', text)
        self.assertEqual(document.paragraphs[1].style.name, 'Heading 2')
        with self.assertRaises(FileExistsError):
            rt.science_export_manuscript(self.workspace, 'draft.md', 'draft.docx')

    @unittest.skipUnless(importlib.util.find_spec('docx') and importlib.util.find_spec('PIL'), 'DOCX/image dependencies unavailable')
    def test_docx_tables_figures_links_and_emphasis(self):
        from docx import Document
        from PIL import Image
        from zipfile import ZipFile
        (self.root / 'figures').mkdir()
        Image.new('RGB', (60, 40), 'blue').save(self.root / 'figures' / 'plot.png')
        (self.root / 'chapters').mkdir()
        manuscript = '# Study\n\n**Strong** and *careful* findings. [Source](https://doi.org/10.1038/nature12373)\n\n| Group | Effect |\n| --- | ---: |\n| A | **2.0** |\n| B | 1.0 |\n\n![Figure 1. Measured effects.](../figures/plot.png)\n'
        (self.root / 'chapters' / 'draft.md').write_text(manuscript, encoding='utf-8')
        result = rt.science_export_manuscript(self.workspace, 'chapters/draft.md', 'draft.docx')
        document = Document(result['path'])
        self.assertEqual(result['tables'], 1)
        self.assertEqual(document.tables[0].cell(1, 1).text, '2.0')
        self.assertEqual(len(document.tables[0].rows), 3)
        self.assertEqual(result['figures'], 1)
        self.assertTrue(any(r.bold and r.text == 'Strong' for p in document.paragraphs for r in p.runs))
        self.assertTrue(any(r.italic and r.text == 'careful' for p in document.paragraphs for r in p.runs))
        self.assertIn('Figure 1. Measured effects.', [p.text for p in document.paragraphs])
        with ZipFile(result['path']) as package:
            self.assertIn('https://doi.org/10.1038/nature12373', package.read('word/_rels/document.xml.rels').decode())
            self.assertIn('w:hyperlink', package.read('word/document.xml').decode())
            self.assertTrue(any(name.startswith('word/media/') for name in package.namelist()))

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'python-docx unavailable')
    def test_export_rejects_remote_and_escaping_figures(self):
        for source in ['https://example.com/figure.png', '../outside.png', 'missing.png']:
            (self.root / 'draft.md').write_text(f'![figure]({source})', encoding='utf-8')
            with self.assertRaises((ValueError, FileNotFoundError)):
                rt.science_export_manuscript(self.workspace, 'draft.md', 'bad.docx')
            self.assertFalse((self.root / 'bad.docx').exists())

    def test_manuscript_findings_are_structural(self):
        text = '# Study\n\n## Abstract\n\n## Results\n\nTODO Add evidence [2]. DOI:???\n\n![Figure](missing.png)\n\n## References\n\n[1] A supplied reference\n'
        (self.root / 'draft.md').write_text(text, encoding='utf-8')
        result = rt.science_check_manuscript(self.workspace, 'draft.md')
        kinds = [item['kind'] for item in result['findings']]
        self.assertIn('empty_section', kinds)
        self.assertIn('unresolved_placeholder', kinds)
        self.assertIn('malformed_doi', kinds)
        self.assertIn('missing_or_invalid_figure', kinds)
        self.assertEqual([f['reference'] for f in result['findings'] if f['kind'] == 'unresolved_numbered_reference'], ['2'])

    def test_manuscript_code_and_resolved_numbered_references(self):
        text = '# Study\n\n## Results\n\nReported result [1]. [Paper](https://doi.org/10.1038/nature12373)\n\n```python\nTODO = "[2]"\n```\n\n## References\n\n[1] Supplied bibliographic entry\n'
        (self.root / 'draft.md').write_text(text, encoding='utf-8')
        result = rt.science_check_manuscript(self.workspace, 'draft.md')
        self.assertEqual(result['findings'], [])
        self.assertEqual(result['doi_candidates'][0]['doi'], '10.1038/nature12373')
        self.assertFalse(result['doi_candidates'][0]['resolved_online'])

    def test_parentheses_in_doi_and_leading_numeric_citation(self):
        text = '# Study\n\n[3] supports this claim.\n\n[Paper](https://doi.org/10.1000/example(123))\n'
        (self.root / 'draft.md').write_text(text, encoding='utf-8')
        result = rt.science_check_manuscript(self.workspace, 'draft.md')
        self.assertEqual(result['doi_candidates'][0]['doi'], '10.1000/example(123)')
        self.assertIn('3', [f['reference'] for f in result['findings'] if f['kind'] == 'unresolved_numbered_reference'])

    def test_bibliography_response_compact_json_file_keeps_detail(self):
        record = dict(RECORD, abstract='Detailed abstract' * 1000, reference=[{'DOI': '10.1000/source'}])
        with patch.object(rt, '_crossref', return_value=record):
            result = rt.science_export_bibliography(self.workspace, ['10.1038/nature12373'], 'refs.json', 'json')
        self.assertNotIn('abstract', result['records'][0])
        self.assertNotIn('references', result['records'][0])
        self.assertEqual(result['records'][0]['title'], RECORD['title'][0])
        saved = json.loads((self.root / 'refs.json').read_text())
        self.assertEqual(saved[0]['abstract'], record['abstract'])
        self.assertEqual(saved[0]['references'], record['reference'])

    def test_pdf_response_budget_and_exact_continuation(self):
        class Page:
            def get_textpage(self): return self
            def get_text_range(self): return 'x' * 100000
            def close(self): pass
        class Document:
            def __init__(self): self.read = []
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def __len__(self): return 5
            def __getitem__(self, index):
                self.read.append(index + 1)
                return Page()
        document = Document()
        with patch.object(rt, '_pdf', return_value=document):
            first = rt.science_pdf_pages(self.workspace, 'paper.pdf', [5, 3, 1], 100000)
        self.assertEqual(document.read, [5, 3])
        self.assertEqual(first['returned_characters'], 200000)
        self.assertEqual(first['pages_read'], 2)
        self.assertEqual(first['remaining_requested_pages'], [1])
        self.assertEqual(first['next_pages'], [1])
        self.assertEqual(first['remaining_pages'], 3)
        with patch.object(rt, '_pdf', return_value=Document()):
            second = rt.science_pdf_pages(self.workspace, 'paper.pdf', first['next_pages'], 100000)
        self.assertEqual(second['next_pages'], [])
        self.assertEqual([p['page'] for p in second['pages']], [1])

    def test_whitespace_quote_is_rejected(self):
        for quote in ['   ', '\n\t']:
            entry = {'claim': 'A claim', 'source': 'paper.pdf', 'page': 1, 'quote': quote}
            with self.assertRaises(ValueError):
                rt.science_save_evidence(self.workspace, [entry], 'invalid.csv')
        self.assertFalse((self.root / 'invalid.csv').exists())

    def test_tools_register_and_have_unique_names(self):
        class Fake:
            def __init__(self): self.functions = []
            def tool(self): return lambda f: self.functions.append(f)
        host = Fake()
        rt.register_tools(host)
        self.assertEqual(len(host.functions), 12)
        self.assertEqual(len({f.__name__ for f in host.functions}), 12)


if __name__ == '__main__':
    unittest.main()
